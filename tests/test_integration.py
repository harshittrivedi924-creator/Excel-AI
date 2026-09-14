import openpyxl
import pytest

from backend.copilot import ExcelCopilot


@pytest.fixture
def sample_workbook(tmp_path):
    """Create a sample Excel workbook with test data."""
    filepath = tmp_path / "test_data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws["A1"] = "Name"
    ws["B1"] = "Age"
    ws["C1"] = "Salary"
    ws["D1"] = "Revenue"

    rows = [
        ("Rahul", 25, 50000, 1000),
        ("Amit", 30, 70000, 2000),
        ("Priya", 28, 60000, 3000),
        ("Sunil", 35, 80000, 4000),
        ("Neha", 22, 45000, 5000),
    ]
    for i, (name, age, salary, revenue) in enumerate(rows, start=2):
        ws[f"A{i}"] = name
        ws[f"B{i}"] = age
        ws[f"C{i}"] = salary
        ws[f"D{i}"] = revenue

    wb.save(filepath)
    return str(filepath)


class TestEndToEnd:
    def test_sum_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("Column D ka total karo")
        assert result["success"] is True
        assert result["result"] == 15000
        assert "Done" in result["message"]

    def test_average_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("average of column B")
        assert result["success"] is True
        assert result["result"] == 28.0

    def test_max_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("maximum of column B")
        assert result["success"] is True
        assert result["result"] == 35

    def test_min_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("minimum of column B")
        assert result["success"] is True
        assert result["result"] == 22

    def test_count_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("count of column B")
        assert result["success"] is True
        assert result["result"] == 5

    def test_arithmetic_cells(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("B2 + C2")
        assert result["success"] is True
        assert result["result"] == 50025

    def test_multiply_cells(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("B2 aur C2 multiply karke D8 mein daalo")
        assert result["success"] is True
        assert result["destination"] == "D8"


class TestWriteToExcel:
    def test_write_sum_to_destination(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("Column D ka total karo aur D21 mein daal do")
        assert result["success"] is True
        assert result["destination"] == "D21"

        # Verify the formula was written
        copilot.excel.save()
        from backend.excel.handler import ExcelHandler

        handler = ExcelHandler(sample_workbook)
        formula = handler.get_formula("D21")
        assert formula is not None
        assert "SUM" in formula

    def test_detect_destination_overwrite(self, sample_workbook):
        """If the destination cell has data and overwrite not allowed, error."""
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("Column D ka total karo aur A2 mein daal do")
        # A2 contains 'Rahul', so this should fail without overwrite permission
        assert result["success"] is False
        assert "contains data" in result["message"]

    def test_overwrite_allowed(self, sample_workbook):
        """With allow_overwrite=True, the operation should succeed."""
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command(
            "Column D ka total karo aur A2 mein daal do",
            allow_overwrite=True,
        )
        assert result["success"] is True


class TestErrors:
    def test_invalid_column(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("sum of column Z")
        assert result["success"] is False
        assert "Column Z" in result["message"]

    def test_empty_range(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("sum of column Z")
        assert result["success"] is False

    def test_unsupported_operation(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("vlookup of column D")
        assert result["success"] is False

    def test_unresolved_destination_rejected(self, sample_workbook):
        """A write-intent marker with no valid destination must not be
        silently ignored (the result was previously reported without writing)."""
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("Column D ka total karo aur ABC mein daalo")
        assert result["success"] is False
        assert "destination" in result["message"]

    def test_multiletter_column_named_in_error(self, sample_workbook):
        """'Column XYZ ...' should fail naming the column, not a keyword like
        'total'."""
        copilot = ExcelCopilot(sample_workbook)
        result = copilot.process_command("Column XYZ ka total karo")
        assert result["success"] is False
        assert "Column XYZ" in result["message"]

    def test_divide_by_zero(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        # Make B cell zero
        copilot.excel.set_cell_value("B2", 0)
        result = copilot.process_command("divide B2 by 0")
        assert result["success"] is False


class TestHistory:
    def test_history_recorded(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        copilot.process_command("Column D ka total karo")
        assert len(copilot.history) == 1
        assert copilot.history[0].result == 15000

    def test_multiple_history(self, sample_workbook):
        copilot = ExcelCopilot(sample_workbook)
        copilot.process_command("Column D ka total karo")
        copilot.process_command("average of column B")
        assert len(copilot.history) == 2


class TestInteractiveOverwrite:
    """Issue #1: overwrite confirmation in CLI interactive mode."""

    def test_accept_overwrite_prompt(self, sample_workbook, monkeypatch):
        copilot = ExcelCopilot(sample_workbook)
        monkeypatch.setattr("builtins.input", lambda prompt: "yes")
        result = copilot.process_interactive("Column D ka total karo aur A2 mein daal do")
        assert result["success"] is True
        assert result["destination"] == "A2"

        copilot.excel.save()
        from backend.excel.handler import ExcelHandler

        handler = ExcelHandler(sample_workbook)
        assert handler.get_formula("A2") == "SUM(D2:D6)"

    def test_decline_overwrite_prompt(self, sample_workbook, monkeypatch):
        copilot = ExcelCopilot(sample_workbook)
        monkeypatch.setattr("builtins.input", lambda prompt: "no")
        result = copilot.process_interactive("Column D ka total karo aur A2 mein daal do")
        assert result["success"] is False
        assert "already contains data" in result["message"]
        assert copilot.excel.get_cell_value("A2") == "Rahul"

    def test_no_prompt_when_overwrite_not_needed(self, sample_workbook, monkeypatch):
        copilot = ExcelCopilot(sample_workbook)
        monkeypatch.setattr(
            "builtins.input",
            lambda prompt: (_ for _ in ()).throw(
                AssertionError("input() should not have been called")
            ),
        )
        result = copilot.process_interactive("Column D ka total karo")
        assert result["success"] is True
        assert result["result"] == 15000

    def test_interrupt_at_prompt_returns_none(self, sample_workbook, monkeypatch):
        copilot = ExcelCopilot(sample_workbook)
        monkeypatch.setattr("builtins.input", lambda prompt: (_ for _ in ()).throw(EOFError))
        result = copilot.process_interactive("Column D ka total karo aur A2 mein daal do")
        assert result is None
