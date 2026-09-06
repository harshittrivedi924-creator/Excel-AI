import openpyxl
import pytest
from datetime import datetime

from backend.copilot import ExcelCopilot


@pytest.fixture
def conditional_workbook(tmp_path):
    path = tmp_path / "cond.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Product"
    ws["B1"] = "Revenue"
    for i, (p, r) in enumerate(
        [("Pen", 100), ("Pencil", 200), ("Pen", 150), ("Eraser", 300)],
        start=2,
    ):
        ws[f"A{i}"] = p
        ws[f"B{i}"] = r
    wb.save(path)
    return path


class TestConditional:
    def test_sumif_text_criteria(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command("Column B ka sum karo jahan Product Pen hai")
        assert result["success"] is True
        assert result["result"] == 250  # 100 + 150

    def test_sumif_symbol_criteria(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command("Column B ka sum karo jahan B > 120")
        assert result["success"] is True
        assert result["result"] == 650  # 200 + 150 + 300

    def test_countif(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command("Count karo jahan Product Pen hai")
        assert result["success"] is True
        assert result["result"] == 2

    def test_averageif(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command("Revenue ka average karo jahan A Pen hai")
        assert result["success"] is True
        assert result["result"] == 125.0

    def test_multiple_conditions(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command(
            "Column B ka sum karo jahan Product Pen hai aur B 120 se zyada"
        )
        assert result["success"] is True
        assert result["result"] == 150

    def test_sumif_writes_formula(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command(
            "Column B ka sum karo jahan Product Pen hai aur C21 mein daal do"
        )
        assert result["success"] is True
        assert result["destination"] == "C21"
        formula = copilot.excel.get_formula("C21")
        assert formula.startswith("SUMIF")

    def test_sumif_no_match(self, conditional_workbook):
        copilot = ExcelCopilot(conditional_workbook)
        result = copilot.process_command(
            "Column B ka sum karo jahan Product Marker hai"
        )
        assert result["success"] is True
        assert result["result"] == 0


@pytest.fixture
def cleaning_workbook(tmp_path):
    path = tmp_path / "clean.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Joining Date"
    ws["C1"] = "Score"
    ws["A2"] = "rahul kumar"
    ws["A3"] = "AMIT"
    ws["B2"] = "15-02-2024"
    ws["C2"] = 100
    ws["C3"] = "not-a-number"
    wb.save(path)
    return path


class TestCleaning:
    def test_standardize_names(self, cleaning_workbook):
        copilot = ExcelCopilot(cleaning_workbook)
        result = copilot.process_command("Names ko standardize karo")
        assert result["success"] is True
        assert result["result"]["converted"] == 2

        copilot.excel.save()
        wb = openpyxl.load_workbook(cleaning_workbook)
        ws = wb["Data"]
        assert ws["A2"].value == "Rahul Kumar"
        assert ws["A3"].value == "Amit"

    def test_standardize_dates(self, cleaning_workbook):
        copilot = ExcelCopilot(cleaning_workbook)
        result = copilot.process_command("Dates ko standardize karo")
        assert result["success"] is True
        assert result["result"]["converted"] == 1

        copilot.excel.save()
        wb = openpyxl.load_workbook(cleaning_workbook)
        assert wb["Data"]["B2"].value == "15-02-2024"

    def test_standardize_dates_datetime_cell(self, tmp_path):
        path = tmp_path / "dates.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "When"
        ws["A2"] = datetime(2024, 3, 1)
        wb.save(path)

        copilot = ExcelCopilot(path)
        result = copilot.process_command("Dates ko standardize karo")
        assert result["result"]["converted"] == 1
        copilot.excel.save()
        assert openpyxl.load_workbook(path).active["A2"].value == "01-03-2024"

    def test_detect_invalid(self, cleaning_workbook):
        copilot = ExcelCopilot(cleaning_workbook)
        result = copilot.process_command("Invalid values detect karo")
        assert result["success"] is True
        assert result["written_count"] >= 1

    def test_dedupe_includes_preview(self, cleaning_workbook):
        wb = openpyxl.load_workbook(cleaning_workbook)
        ws = wb["Data"]
        ws["A4"] = "rahul kumar"
        ws["B4"] = "15-02-2024"
        ws["C4"] = 100
        wb.save(cleaning_workbook)

        copilot = ExcelCopilot(cleaning_workbook)
        result = copilot.process_command("Duplicate rows hata do")
        assert result["success"] is True
        assert result["result"]["preview"]  # removed rows reported


class TestAnomalies:
    def test_analyze_reports_anomaly(self, tmp_path):
        path = tmp_path / "anomaly.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Val"
        for i, v in enumerate([10, 11, 12, 9, 10, 150], start=2):
            ws[f"A{i}"] = v
        wb.save(path)

        copilot = ExcelCopilot(path)
        result = copilot.process_command("Is workbook ka analysis kar do")
        assert result["success"] is True
        lines = "\n".join(result["result"]["lines"])
        assert "Anomaly" in lines