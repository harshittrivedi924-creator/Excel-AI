import openpyxl
import pytest

from backend.excel.handler import ExcelHandler


@pytest.fixture
def sample_workbook(tmp_path):
    """Create a sample Excel workbook with test data."""
    filepath = tmp_path / "test_data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Header row
    ws["A1"] = "Name"
    ws["B1"] = "Age"
    ws["C1"] = "Salary"
    ws["D1"] = "Revenue"

    # Data rows
    rows = [
        ("Rahul", 25, 50000, 1000),
        ("Amit", 30, 70000, 2000),
        ("Priya", 28, 60000, 3000),
        ("Sunil", 35, 80000, 4000),
        ("Neha", 22, 45000, 5000),
    ]
    for i, (name, age, salary, revenue) in enumerate(rows, start=2):
        ws[f"A{i}"] = name
        ws[f"B{i}"] = age
        ws[f"C{i}"] = salary
        ws[f"D{i}"] = revenue

    wb.save(filepath)
    return str(filepath)


class TestReadWorkbook:
    def test_load_workbook(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        assert handler.workbook is not None

    def test_get_sheet_names(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        assert "Sheet1" in handler.get_sheet_names()

    def test_get_cell_value(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        assert handler.get_cell_value("A2") == "Rahul"
        assert handler.get_cell_value("B2") == 25


class TestReadColumn:
    def test_get_used_range(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        used_range = handler.get_used_range("D")
        assert used_range == "D2:D6"

    def test_get_column_values(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        values, max_row, min_row = handler.get_column_values("D")
        assert values == [1000, 2000, 3000, 4000, 5000]
        assert max_row == 6
        assert min_row == 2

    def test_get_column_values_no_header(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        values, max_row, min_row = handler.get_column_values("B")
        assert values == [25, 30, 28, 35, 22]
        assert max_row == 6

    def test_get_range_values(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        values = handler.get_range_values("B2:B4")
        assert values == [25, 30, 28]


class TestWriteFormulas:
    def test_write_formula(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        formula = handler.write_formula("D7", "=SUM(D2:D6)")
        assert formula == "=SUM(D2:D6)"

    def test_write_formula_adds_equal_sign(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        formula = handler.write_formula("D7", "SUM(D2:D6)")
        assert formula == "=SUM(D2:D6)"

    def test_get_formula(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        handler.write_formula("D7", "=SUM(D2:D6)")
        formula = handler.get_formula("D7")
        assert formula == "SUM(D2:D6)"

    def test_cell_has_data(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        assert handler.cell_has_data("A2")
        assert not handler.cell_has_data("A100")

    def test_set_cell_value(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        handler.set_cell_value("E1", "Net Profit")
        assert handler.get_cell_value("E1") == "Net Profit"


class TestSave:
    def test_save_workbook(self, sample_workbook):
        handler = ExcelHandler(sample_workbook)
        handler.set_cell_value("Z1", 42)
        handler.save()
        # Reload and verify
        handler2 = ExcelHandler(sample_workbook)
        assert handler2.get_cell_value("Z1") == 42
