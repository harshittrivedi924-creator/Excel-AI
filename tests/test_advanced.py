import openpyxl
import pytest

from backend.copilot import ExcelCopilot


@pytest.fixture
def finance_workbook(tmp_path):
    """Create a workbook with named columns for advanced tests."""
    filepath = tmp_path / "finance.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Finance"

    ws["A1"] = "Month"
    ws["B1"] = "Revenue"
    ws["C1"] = "Expense"
    ws["D1"] = "Profit"

    rows = [
        ("Jan", 1000, 600),
        ("Feb", 1200, 700),
        ("Mar", 1500, 800),
        ("Apr", 1100, 500),
        ("May", 1800, 900),
    ]
    for i, (month, rev, exp) in enumerate(rows, start=2):
        ws[f"A{i}"] = month
        ws[f"B{i}"] = rev
        ws[f"C{i}"] = exp

    wb.save(filepath)
    return str(filepath)


@pytest.fixture
def month_workbook(tmp_path):
    """Workbook with monthly sales columns."""
    filepath = tmp_path / "months.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "Product"
    ws["B1"] = "January"
    ws["C1"] = "February"
    ws["E1"] = "March"

    ws["A2"] = "P1"
    ws["B2"] = 100
    ws["C2"] = 120
    ws["A3"] = "P2"
    ws["B3"] = 200
    ws["C3"] = 250
    ws["A4"] = "P3"
    ws["B4"] = 300
    ws["C4"] = 330

    wb.save(filepath)
    return str(filepath)


@pytest.fixture
def bonus_workbook(tmp_path):
    """Workbook whose numeric column uses a non-vocabulary header name."""
    filepath = tmp_path / "bonus.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Employee"
    ws["B1"] = "Bonus"
    ws["A2"] = "R1"
    ws["B2"] = 100
    ws["A3"] = "R2"
    ws["B3"] = 200
    ws["A4"] = "R3"
    ws["B4"] = 300
    wb.save(filepath)
    return str(filepath)


@pytest.fixture
def zero_workbook(tmp_path):
    """Workbook with a numeric column whose total is zero."""
    filepath = tmp_path / "zero.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Month"
    ws["B1"] = "Zero"
    for i in range(2, 5):
        ws[f"B{i}"] = 0
    wb.save(filepath)
    return str(filepath)


class TestNamedColumns:
    def test_sum_named_column(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue ka total karo")
        assert result["success"] is True
        assert result["result"] == 1000 + 1200 + 1500 + 1100 + 1800

    def test_average_named_column(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Expense ka average nikal do")
        assert result["success"] is True
        assert result["result"] == pytest.approx((600 + 700 + 800 + 500 + 900) / 5)

    def test_max_named_column(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue ka maximum batao")
        assert result["success"] is True
        assert result["result"] == 1800


class TestProfitCalculation:
    def test_revenue_minus_expense(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue minus expense karke profit nikalo")
        assert result["success"] is True
        assert result["written_count"] == 5

        # Verify formulas written to column D
        copilot.excel.save()
        from backend.excel.handler import ExcelHandler

        handler = ExcelHandler(finance_workbook)
        assert handler.get_formula("D2") == "B2-C2"
        assert handler.get_formula("D3") == "B3-C3"


class TestGrowth:
    def test_growth_between_months(self, month_workbook):
        copilot = ExcelCopilot(month_workbook)
        result = copilot.process_command("February ki sales January se kitni badhi")
        assert result["success"] is True
        jan_total = 100 + 200 + 300
        feb_total = 120 + 250 + 330
        expected = (feb_total - jan_total) / jan_total * 100
        assert result["result"] == pytest.approx(expected)
        assert "Growth from January to February" in result["message"]
        assert "16.67%" in result["message"]

    def test_growth_cells(self, month_workbook):
        copilot = ExcelCopilot(month_workbook)
        result = copilot.process_command("C2 se B2 kitni badhi")
        assert result["success"] is True
        assert result["result"] == pytest.approx(20.0)
        assert "Growth from B2 to C2" in result["message"]
        assert "20%" in result["message"]

    def test_growth_month_per_row_gives_hint(self, finance_workbook):
        """When months are stored as rows (Month column), growth between months
        cannot be computed and should explain why rather than crash."""
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("February vs January growth")
        assert result["success"] is False
        assert "column headers" in result["message"]


class TestDifference:
    def test_difference_cells(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("B2 aur C2 ka difference batao")
        assert result["success"] is True
        assert result["result"] == 400


class TestWholeSheet:
    def test_whole_sheet_sum(self, month_workbook):
        copilot = ExcelCopilot(month_workbook)
        result = copilot.process_command("Sabka sum kar do")
        assert result["success"] is True
        expected = 100 + 200 + 300 + 120 + 250 + 330
        assert result["result"] == expected


class TestColumnArithmetic:
    def test_multiply_columns_into_destination(self, month_workbook):
        copilot = ExcelCopilot(month_workbook)
        result = copilot.process_command("B aur C ko multiply karke F mein daal do")
        assert result["success"] is True
        assert result["written_count"] >= 3

        copilot.excel.save()
        from backend.excel.handler import ExcelHandler

        handler = ExcelHandler(month_workbook)
        assert handler.get_formula("F2") == "B2*C2"


class TestCellPlusNumber:
    def test_cell_plus_number(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("B2 + 100")
        assert result["success"] is True
        assert result["result"] == 1100

    def test_number_plus_cell(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("100 + C2")
        assert result["success"] is True
        assert result["result"] == 700


class TestPercentage:
    def test_percentage_requires_column(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Profit ka percentage calculate karo")
        # Profit column has no values yet (empty), expect a helpful error
        assert result["success"] is False
        assert "numeric" in result["message"] or "column" in result["message"]

    def test_percentage_no_output_reports_row_shares(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue ka percentage karo")
        # Must not silently report the column total (6600) as the answer;
        # instead each value's share of the total is reported.
        assert result["success"] is True
        assert result["result"] == [
            round(v / 6600 * 100, 2) for v in [1000, 1200, 1500, 1100, 1800]
        ]
        assert "6600" in result["message"]
        assert "15.15%" in result["message"]

    def test_percentage_with_destination_writes_formulas(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue ka percentage E2 mein daalo")
        assert result["success"] is True
        assert result["written_count"] == 5
        assert result["result"] == 6600
        copilot.excel.save()
        wb = openpyxl.load_workbook(finance_workbook)
        ws = wb["Finance"]
        assert ws["E2"].value == "=B2/6600*100"
        assert ws["E6"].value == "=B6/6600*100"

    def test_percentage_destination_overwrite_blocked(self, finance_workbook):
        wb = openpyxl.load_workbook(finance_workbook)
        wb["Finance"]["E2"] = 42
        wb.save(finance_workbook)
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Revenue ka percentage E2 mein daalo")
        assert result["success"] is False
        assert "already contains data" in result["message"]

    def test_percentage_arbitrary_header(self, bonus_workbook):
        copilot = ExcelCopilot(bonus_workbook)
        result = copilot.process_command("Bonus ka percentage karo")
        assert result["success"] is True
        assert result["result"] == [round(v / 600 * 100, 2) for v in [100, 200, 300]]
        assert "600" in result["message"]

    def test_percentage_letter_column(self, finance_workbook):
        copilot = ExcelCopilot(finance_workbook)
        result = copilot.process_command("Column B ka percentage karo")
        assert result["success"] is True
        assert result["result"] == [
            round(v / 6600 * 100, 2) for v in [1000, 1200, 1500, 1100, 1800]
        ]

    def test_percentage_zero_total_error(self, zero_workbook):
        copilot = ExcelCopilot(zero_workbook)
        result = copilot.process_command("Zero ka percentage karo")
        assert result["success"] is False
        assert "0" in result["message"]
        assert "cannot be calculated" in result["message"]
