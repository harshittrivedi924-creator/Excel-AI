import io

import openpyxl
import pytest


@pytest.fixture
def api_client(tmp_path):
    from backend.api.app import create_app

    filepath = tmp_path / "web.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Salary"
    for i, (n, s) in enumerate([("Amit", 70000), ("Rahul", 50000)], start=2):
        ws[f"A{i}"] = n
        ws[f"B{i}"] = s
    wb.save(filepath)

    app = create_app(initial_filepath=str(filepath))
    app.config["TESTING"] = True
    return app.test_client()


class TestWebAPI:
    def test_index(self, api_client):
        res = api_client.get("/")
        assert res.status_code == 200
        assert b"Excel AI Copilot" in res.data

    def test_workbook_info(self, api_client):
        res = api_client.get("/api/workbook")
        assert res.status_code == 200
        data = res.get_json()
        assert data["loaded"] is True
        assert data["sheets"] == ["Sheet"]

    def test_command(self, api_client):
        res = api_client.post("/api/command", json={"command": "Column B ka total karo"})
        assert res.status_code == 200
        data = res.get_json()
        assert data["success"] is True
        assert "120000" in data["message"]

    def test_command_overwrite_flag(self, api_client):
        # A2 contains "Amit"; writing a total there needs confirmation
        res = api_client.post(
            "/api/command",
            json={"command": "Column B ka total karo aur A2 mein daal do"},
        )
        data = res.get_json()
        assert data["success"] is False
        assert data["overwrite_needed"] is True

        res = api_client.post(
            "/api/command",
            json={"command": "Column B ka total karo aur A2 mein daal do", "allow_overwrite": True},
        )
        data = res.get_json()
        assert data["success"] is True

    def test_command_without_workbook(self, tmp_path):
        from backend.api.app import create_app

        app = create_app()
        app.config["TESTING"] = True
        client = app.test_client()
        res = client.post("/api/command", json={"command": "sum karo"})
        assert res.status_code == 400

    def test_upload(self, api_client, tmp_path):
        fp = tmp_path / "upload.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "X"
        ws["B1"] = "Y"
        ws["B2"] = 42
        wb.save(fp)

        with open(fp, "rb") as f:
            res = api_client.post(
                "/api/upload",
                data={"file": (io.BytesIO(f.read()), "upload.xlsx")},
                content_type="multipart/form-data",
            )
        assert res.status_code == 200
        data = res.get_json()
        assert data["loaded"] is True
