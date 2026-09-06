import os

from flask import Flask, jsonify, request, send_from_directory

from backend.copilot import ExcelCopilot

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
FRONTEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "..", "frontend")


def create_app(initial_filepath=None):
    app = Flask(__name__)
    copilot = ExcelCopilot()
    active_file = {"path": None}

    if initial_filepath and os.path.exists(initial_filepath):
        copilot.load_workbook(initial_filepath)
        active_file["path"] = os.path.abspath(initial_filepath)

    @app.route("/")
    def index():
        return send_from_directory(FRONTEND_DIR, "index.html")

    @app.route("/api/workbook")
    def workbook_info():
        if not copilot.excel.workbook:
            return jsonify({"loaded": False, "sheets": [], "rows": [], "file": None})

        try:
            from openpyxl.utils import get_column_letter
            data = []
            sheet = copilot.excel.get_sheet()
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, 12),
                max_col=min(sheet.max_column, 12),
                values_only=True,
            ):
                data.append(["" if v is None else str(v) for v in row])
            return jsonify({
                "loaded": True,
                "sheets": copilot.excel.get_sheet_names(),
                "active_sheet": copilot.excel.get_sheet().title,
                "rows": data,
                "file": os.path.basename(active_file["path"]) if active_file["path"] else None,
            })
        except Exception as e:
            return jsonify({"loaded": False, "error": str(e), "sheets": [], "rows": []})

    @app.route("/api/upload", methods=["POST"])
    def upload():
        file = request.files.get("file")
        if not file:
            return jsonify({"error": "No file uploaded"}), 400

        os.makedirs(UPLOAD_DIR, exist_ok=True)
        path = os.path.join(UPLOAD_DIR, file.filename)
        file.save(path)
        try:
            copilot.load_workbook(path)
            active_file["path"] = path
            return workbook_info()
        except Exception as e:
            return jsonify({"loaded": False, "error": f"Couldn't open file: {e}"}), 400

    @app.route("/api/command", methods=["POST"])
    def command():
        payload = request.get_json(force=True, silent=True) or {}
        cmd = payload.get("command", "")
        sheet = payload.get("sheet")
        allow_overwrite = bool(payload.get("allow_overwrite", False))

        if not copilot.excel.workbook:
            return jsonify({
                "success": False,
                "message": "No workbook loaded. Upload an Excel file first.",
            }), 400

        result = copilot.process_command(cmd, sheet_name=sheet,
                                         allow_overwrite=allow_overwrite)

        if result["success"] and active_file["path"]:
            try:
                copilot.excel.save()
            except Exception:
                pass

        return jsonify({
            "success": result["success"],
            "message": result["message"],
            "result": result.get("result"),
            "overwrite_needed": (
                not result["success"]
                and result.get("message", "")
                and "contains data" in result["message"]
            ),
        })

    @app.route("/api/history")
    def history():
        items = [{
            "operation": e.instruction.get("operation"),
            "result": e.result,
            "destination": e.destination,
        } for e in copilot.history]
        return jsonify(items)

    return app